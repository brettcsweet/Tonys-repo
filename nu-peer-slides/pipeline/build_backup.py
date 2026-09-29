"""Backup workbook: data behind the three slides, live formulas for every derived number, checks log, data still needed."""
import sys
import numpy as np, pandas as pd, openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = sys.argv[1]
F = "Lato"
def font(**k): return Font(name=F, size=k.pop("size", 10), **k)
HDR_BORDER = Border(bottom=Side(style="thin", color="000000"))
RED = "C8102E"

def header(ws, row, labels, widths=None):
    for i, l in enumerate(labels, 1):
        c = ws.cell(row, i, l); c.font = font(bold=True); c.border = HDR_BORDER
        c.alignment = Alignment(wrap_text=True, vertical="bottom")
    if widths:
        for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w

def title(ws, text, sub=None):
    ws["A1"] = text; ws["A1"].font = font(bold=True, size=14, color=RED)
    if sub: ws["A2"] = sub; ws["A2"].font = font(size=9, color="6B6B70")

wb = Workbook()

# ------------------------------------------------------------------ README
ws = wb.active; ws.title = "README"
title(ws, "Peer slides: backup data and checks", "Interim build on the data inside the uploaded deck. Every derived number below is a live formula.")
lines = [
    ("Status", "Slides are rebranded. Data has NOT yet been refreshed: the network policy of the working session blocks the sources (see Data_needed)."),
    ("Slide 1", "S1_Scorecard: 1,557 institutions from the Oct 2023 College Scorecard release (as embedded in the original chart). Northeastern highlighted."),
    ("Slide 2", "S2_Survey: Inside Higher Ed business-officer survey values 2014-2024 as printed on the original slide. 2014 removed from the slide per request."),
    ("Slide 3", "S3_FY23: FY23 salaries and wages, core operating expenses and students as embedded in the original chart. Northeastern, BU, MIT not yet added."),
    ("Checks", "Checks: every test run on the original numbers, with result. Findings are items to resolve in the refresh."),
    ("Data_needed", "Exact files and hosts required to finish the update."),
]
for i, (a, b) in enumerate(lines, 4):
    ws.cell(i, 1, a).font = font(bold=True); ws.cell(i, 2, b).font = font(); ws.cell(i, 2).alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 14; ws.column_dimensions["B"].width = 120

# ------------------------------------------------------------------ S1
df = pd.read_csv("inputs/slide1_scorecard_oct2023.csv")
df["ctl"] = df.control.map({1.0: "Public", 2.0: "Private"})
df = df.sort_values("y", ascending=False).reset_index(drop=True)
ws = wb.create_sheet("S1_Scorecard")
title(ws, "Slide 1: net price vs median earnings 10 years after entry (College Scorecard, Oct 2023 release)")
header(ws, 4, ["Institution", "City", "State", "Control", "Avg annual net price (NPT4)", "Median earnings, 10th yr after entry (MD_EARN_WNE_P10)",
               "Rank by earnings", "Highlight", "Private net price", "Private earnings"], [46, 20, 7, 9, 16, 22, 10, 18, 14, 14])
r0 = 5
hl = {"Northeastern University": "Northeastern", "Vanderbilt University": "Vanderbilt", "Boston University": "Boston University",
      "Massachusetts Institute of Technology": "MIT"}
n = len(df); r1 = r0 + n - 1
for i, row in df.iterrows():
    r = r0 + i
    ws.cell(r, 1, row["name"]); ws.cell(r, 2, row["city"]); ws.cell(r, 3, row["st"]); ws.cell(r, 4, row["ctl"])
    ws.cell(r, 5, int(row["x"])).number_format = '"$"#,##0'; ws.cell(r, 6, int(row["y"])).number_format = '"$"#,##0'
    ws.cell(r, 7, f"=RANK(F{r},$F${r0}:$F${r1})")
    ws.cell(r, 8, hl.get(row["name"], ""))
    ws.cell(r, 9, f'=IF(D{r}="Private",E{r},"")').number_format = '"$"#,##0'
    ws.cell(r, 10, f'=IF(D{r}="Private",F{r},"")').number_format = '"$"#,##0'
    for c in range(1, 11): ws.cell(r, c).font = font(bold=bool(hl.get(row["name"])), color=(RED if row["name"] == "Northeastern University" else "000000"))
# summary block
sc = 12
ws.cell(4, sc, "Summary (formulas)").font = font(bold=True); ws.cell(4, sc).border = HDR_BORDER; ws.cell(4, sc + 1).border = HDR_BORDER
summ = [
    ("Institutions", f"=COUNTA(A{r0}:A{r1})"),
    ("  Public", f'=COUNTIF(D{r0}:D{r1},"Public")'),
    ("  Private", f'=COUNTIF(D{r0}:D{r1},"Private")'),
    ("Vanderbilt rank (slide title says 53)", f'=INDEX(G{r0}:G{r1},MATCH("Vanderbilt University",A{r0}:A{r1},0))'),
    ("Northeastern rank", f'=INDEX(G{r0}:G{r1},MATCH("Northeastern University",A{r0}:A{r1},0))'),
    ("Boston University rank", f'=INDEX(G{r0}:G{r1},MATCH("Boston University",A{r0}:A{r1},0))'),
    ("MIT rank", f'=INDEX(G{r0}:G{r1},MATCH("Massachusetts Institute of Technology",A{r0}:A{r1},0))'),
    ("Private trendline slope", f"=SLOPE(J{r0}:J{r1},I{r0}:I{r1})"),
    ("Private trendline intercept", f"=INTERCEPT(J{r0}:J{r1},I{r0}:I{r1})"),
    ("Private trendline r", f"=CORREL(I{r0}:I{r1},J{r0}:J{r1})"),
    ("Private trendline R-squared", f"=RSQ(J{r0}:J{r1},I{r0}:I{r1})"),
    ("Northeastern earnings", f'=INDEX(F{r0}:F{r1},MATCH("Northeastern University",A{r0}:A{r1},0))'),
    ("Northeastern net price", f'=INDEX(E{r0}:E{r1},MATCH("Northeastern University",A{r0}:A{r1},0))'),
    ("Northeastern vs private trendline ($)", None),
]
for i, (lab, f) in enumerate(summ, 5):
    ws.cell(i, sc, lab).font = font()
    if f: ws.cell(i, sc + 1, f).font = font(bold=True)
last = 5 + len(summ) - 1
ws.cell(last, sc + 1, f"=M{last-2}-(M{last-5}+M{last-6}*M{last-1})").font = font(bold=True)
ws.cell(last, sc + 1).number_format = '"$"#,##0'
for r in (5, 6, 7, 8, 9, 10, 11): ws.cell(r, sc + 1).number_format = "0"
ws.cell(12, sc + 1).number_format = "0.000"; ws.cell(13, sc + 1).number_format = '"$"#,##0'
ws.cell(14, sc + 1).number_format = "0.000"; ws.cell(15, sc + 1).number_format = "0.000"
ws.cell(16, sc + 1).number_format = '"$"#,##0'; ws.cell(17, sc + 1).number_format = '"$"#,##0'
ws.column_dimensions[get_column_letter(sc)].width = 40; ws.column_dimensions[get_column_letter(sc + 1)].width = 14
ws.freeze_panes = "A5"

# ------------------------------------------------------------------ S2
import json
raw = json.load(open("inputs/slide2_data_orig.json"))
ws = wb.create_sheet("S2_Survey")
title(ws, "Slide 2: '_____ are realistic and aware of the financial challenges confronting my institution'",
      "Percent of business officers answering 4 or 5 on a five-point scale. Values as printed on the original slide.")
header(ws, 4, ["Survey year", "Faculty", "Trustees", "Sr. Administrators", "Gap: Trustees minus Faculty (pts)", "Gap: Sr. Admin minus Faculty (pts)", "On slide?"], [12, 10, 10, 18, 22, 22, 34])
for i, yr in enumerate(range(2014, 2025), 5):
    ws.cell(i, 1, yr)
    for j, k in enumerate("FTS", 2): ws.cell(i, j, raw[k][str(yr)] / 100).number_format = "0%"
    ws.cell(i, 5, f"=(C{i}-B{i})*100").number_format = "0"; ws.cell(i, 6, f"=(D{i}-B{i})*100").number_format = "0"
    ws.cell(i, 7, "Removed per request" if yr == 2014 else "Yes")
    for c in range(1, 8): ws.cell(i, c).font = font()
ws.cell(16, 1, 2025).font = font(); ws.cell(16, 7, "Not yet sourced (see Data_needed)").font = font(color=RED)
ws.cell(17, 1, 2026).font = font(); ws.cell(17, 7, "Newer edition exists (July 2026); not yet sourced").font = font(color=RED)
ws.cell(19, 1, "Averages 2015-2024").font = font(bold=True)
for j, col in enumerate("BCDEF", 2):
    ws.cell(19, j, f"=AVERAGE({col}6:{col}15)").number_format = "0%" if col in "BCD" else "0.0"; ws.cell(19, j).font = font(bold=True)

# ------------------------------------------------------------------ S3
ws3 = openpyxl.load_workbook("inputs/orig_slide3_workbook.xlsx", data_only=True).active
norm = {"U Chicago": "Chicago", "U Penn": "Penn"}
def block(r0):
    d = {}
    for r in range(r0, r0 + 13):
        nm = ws3.cell(r, 13).value
        if nm in (None, "University", "Institution"): continue
        v = [ws3.cell(r, c).value for c in (14, 15, 16)]
        if all(isinstance(x, (int, float)) for x in v): d[norm.get(nm, nm)] = v
    return d
fy23, fy22 = block(94), block(81)
label_sw = {"Rice": "$51K", "Vanderbilt": "$51K", "Northwestern": "$53K", "Princeton": "$110K", "Harvard": "$113K", "Yale": "$158K",
            "Chicago": "$168K", "Duke": "$204K", "Emory": "$269 (K missing)", "Penn": "$241K", "Stanford": "$405K"}
label_st = {"Emory": "16,00 (truncated)"}
HEALTH = {"Chicago", "Emory", "Penn", "Duke", "Stanford"}
ws = wb.create_sheet("S3_FY23")
title(ws, "Slide 3: FY23 salaries & wages per student (inputs from the original chart's embedded workbook)",
      "Derived columns are formulas. FY22 columns are from the same embedded workbook and are used only for the consistency check.")
header(ws, 4, ["Institution", "S&W ($)", "Core operating expenses ($)", "Students", "S&W per student", "Opex per student (x)", "S&W share of opex (y)",
               "Check: x*y - S&W/student", "Label on original slide (S&W/student)", "Health care system included?",
               "FY22 S&W ($)", "FY22 core opex ($)", "S&W growth FY22-23", "Opex growth FY22-23", "S&W share change (pts)"],
       [16, 16, 18, 10, 14, 14, 14, 16, 22, 14, 16, 16, 12, 12, 14])
order = ["Rice", "Vanderbilt", "Northwestern", "Princeton", "Harvard", "Yale", "Chicago", "Duke", "Emory", "Penn", "Stanford"]
for i, nm in enumerate(order, 5):
    sw, ox, st = fy23[nm]; sw2, ox2, _ = fy22[nm]
    vals = [nm, sw, ox, st, f"=B{i}/D{i}", f"=C{i}/D{i}", f"=B{i}/C{i}", f"=F{i}*G{i}-E{i}", label_sw[nm], "Yes" if nm in HEALTH else "No",
            sw2, ox2, f"=B{i}/K{i}-1", f"=C{i}/L{i}-1", f"=(B{i}/C{i}-K{i}/L{i})*100"]
    for j, v in enumerate(vals, 1):
        c = ws.cell(i, j, v); c.font = font()
    for j in (2, 3, 11, 12): ws.cell(i, j).number_format = "#,##0"
    ws.cell(i, 4).number_format = "#,##0"
    for j in (5, 6): ws.cell(i, j).number_format = '"$"#,##0'
    ws.cell(i, 7).number_format = "0.0%"; ws.cell(i, 8).number_format = "0.00"
    ws.cell(i, 13).number_format = "0.0%"; ws.cell(i, 14).number_format = "0.0%"; ws.cell(i, 15).number_format = "0.0"
end = 4 + len(order)
ws.cell(end + 2, 1, "Median").font = font(bold=True)
for col in "MNO":
    ws[f"{col}{end+2}"] = f"=MEDIAN({col}5:{col}{end})"; ws[f"{col}{end+2}"].font = font(bold=True)
ws[f"M{end+2}"].number_format = "0.0%"; ws[f"N{end+2}"].number_format = "0.0%"; ws[f"O{end+2}"].number_format = "0.0"
ws.cell(end + 3, 1, "Max / min S&W per student").font = font(bold=True)
ws[f"E{end+3}"] = f"=MAX(E5:E{end})/MIN(E5:E{end})"; ws[f"E{end+3}"].number_format = '0.0"x"'; ws[f"E{end+3}"].font = font(bold=True)
ws.cell(end + 4, 1, "Institutions with S&W share falling FY22-23").font = font(bold=True)
ws[f"E{end+4}"] = f'=COUNTIF(O5:O{end},"<0")&" of "&COUNT(O5:O{end})'; ws[f"E{end+4}"].font = font(bold=True)

# ------------------------------------------------------------------ Checks
ws = wb.create_sheet("Checks")
title(ws, "Checks run on the original numbers, and disconnects found")
header(ws, 3, ["#", "Slide", "Check / observation", "Result", "Status"], [5, 7, 70, 90, 12])
checks = [
    (1, "1", "587 public + 970 private = 1,557 institutions", "Confirmed from the embedded data (see S1_Scorecard summary).", "Pass"),
    (1, "1", "Vanderbilt ranks 53rd of 1,557 by median earnings (title)", "Confirmed; no ties at $84,415.", "Pass"),
    (1, "1", "Plotted 'private college trendline' = OLS fit on private institutions only", "Reproduced: slope 1.035, intercept $29,753, r = 0.49, R-squared = 0.24. Weak fit: 'better deal / worse deal' is directional only.", "Pass / caveat"),
    (1, "1", "Northeastern position (Oct 2023 vintage)", "Rank 38 of 1,557 (34 of 970 private); earnings $88,842 at net price $34,255; about $23.6K above the private trendline. BU rank 75; MIT rank 2.", "Info"),
    (1, "1", "Title/axis say '10 years after graduation'; footnote 1 says '10th year after enrollment'", "Field MD_EARN_WNE_P10 is measured from enrollment. Slide now says 'after entry'. For 4-year graduates that is about 6 years after graduation; for co-op/5-year programs about 5.", "Fixed"),
    (1, "1", "Price and earnings are on different bases", "Net price is a recent-year figure for Title IV aided students; earnings are for aided students who enrolled about a decade earlier and are working, not enrolled. Excludes students who go on to graduate school, so schools that feed graduate programs are understated.", "Caveat"),
    (1, "1", "Institutions no longer separate", "University of the Sciences merged into Saint Joseph's in 2022 but is still in the Oct 2023 file. 'Mills College at Northeastern University' (CA) is a separate row; the highlight uses Northeastern University (Boston, MA).", "Open (refresh)"),
    (1, "1", "MIT net price $5,084 is the lowest in the file by a wide margin", "Value is as published in that release; spot-check on the refreshed file before showing MIT as a benchmark.", "Open (refresh)"),
    (1, "1", "Vanderbilt-internal speaker notes ('VU Factbook Master Slide', preparer name)", "Not carried into the new deck.", "Fixed"),
    (2, "2", "Values read from the original slide (30 points, 2015-2024)", "See S2_Survey. Dot positions on the original were hand-placed shapes; fitted to their labels they are off by 1.2 points on average, 3.6 at most, so no reading changes.", "Pass"),
    (2, "2", "Faculty vs Trustees gap over time", "46 pts in 2015 (33% vs 79%), 47 pts in 2024 (44% vs 91%): persistent, not widening. Faculty vs Sr. Admin: 55 pts in 2015 to 51 pts in 2024. Since 2015: Faculty +11 pts (33% to 44%), Trustees +12 (79% to 91%), Sr. Administrators +7 (88% to 95%).", "Info"),
    (2, "2", "Who is being asked", "Respondents are business officers rating other groups (and Sr. Administrators includes their own peers). Slide now says so. It is not faculty's own view of the finances.", "Fixed"),
    (2, "2", "Source line said 2013-2024 and a note covered 2013 trustees", "Both refer to years no longer shown; updated to 2015-2024 and note replaced.", "Fixed"),
    (2, "2", "2025 and 2026 editions", "Inside Higher Ed published the 2025 survey (July 17, 2025) and a 2026 edition (July 15, 2026). Both reports are titled 'Chief Business Officers'. The search summaries reachable from the working session did not show this item, so whether it was asked is unconfirmed. No values were plotted for either year.", "Open"),
    (3, "3", "Identity: S&W per student = x * y for every institution", "Column H of S3_FY23 is zero for all 11 rows; each point sits on its iso-line.", "Pass"),
    (3, "3", "Stanford x-value $915,817 exceeds the original axis maximum $900,000", "Original plotted no institution data at all (only 8 iso-lines); logos were hand-placed pictures, so positions were not data-driven. Rebuilt with real points on a $0-$1,000K axis.", "Fixed"),
    (3, "3", "Label typos on the original", "Iso-line labels '$160', '$200', '$240', '$280', '$320' had no K; Emory '$269' (should round to $270K) and '16,00'.", "Fixed"),
    (3, "3", "Student counts", "Emory (16,000) and Northwestern (23,000) are round numbers, so estimates; others look exact. The refresh should use one IPEDS fall-enrollment definition for all schools.", "Open (refresh)"),
    (3, "3", "FY22 to FY23 consistency (triangulation)", "Core operating expenses rose 15-38% (median 23%) at all 11 schools while S&W rose 7-20% (median 11%); S&W share fell at all 11 (median -4.9 pts). That points to a definitional break between the FY22 and FY23 columns, not an economic change. FY23 vs FY25 must be rebuilt on one definition from source filings, not extended from this sheet.", "Finding"),
    (3, "3", "Copy-forward in older years", "Identical S&W or opex in consecutive years for several schools (e.g., Chicago, Duke, Harvard FY18=FY19; Emory, Northwestern, Rice S&W FY21=FY22). Not displayed on the slide but shows the sheet was hand-assembled.", "Finding"),
    (3, "3", "Comparability of 'S&W per student'", "Health systems (Chicago, Emory, Penn, Duke, Stanford) inflate it and are footnoted. Stanford also includes SLAC; MIT includes Lincoln Laboratory, which has S&W and no students. Add a footnote or a no-health/no-FFRDC view when MIT is added. Total headcount (undergraduate + graduate + professional) also mixes very different cost structures.", "Open (refresh)"),
    (3, "3", "'More cost-efficient than our peers' claim", "S&W per student measures intensity, not efficiency; it cannot support an efficiency claim on its own. Title reworded to a descriptive statement until Northeastern, BU and MIT are added.", "Finding"),
]
for i, (a, b, c, d, e) in enumerate(checks, 4):
    for j, v in enumerate([i - 3, b, c, d, e], 1):
        cell = ws.cell(i, j, v); cell.font = font(); cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(i, 5).font = font(bold=True, color=(RED if e.startswith(("Open", "Finding")) else "000000"))

# ------------------------------------------------------------------ Data_needed
ws = wb.create_sheet("Data_needed")
title(ws, "What is needed to finish the update", "Blocked hosts were denied by the working session's network policy. Alternatively the files can be uploaded to the session.")
header(ws, 4, ["Slide", "Item", "Source", "Host (blocked)", "Status"], [7, 62, 62, 44, 12])
need = [
    ("1", "Latest College Scorecard institution-level file (Most Recent Cohorts - Institution) and data dictionary; confirm earnings cohort year", "collegescorecard.ed.gov/data", "collegescorecard.ed.gov, ed-public-download.scorecard.network", "Blocked"),
    ("1", "IPEDS institutional characteristics for the same UNITIDs (control, Carnegie CCBASIC / CCUGPROF filters)", "nces.ed.gov/ipeds/datacenter", "nces.ed.gov", "Blocked"),
    ("2", "Inside Higher Ed 2025 survey report (July 2025) and 2026 report (July 2026): confirm the 'realistic and aware' item and read faculty / trustee / Sr. Admin values", "insidehighered.com/reports", "www.insidehighered.com", "Blocked"),
    ("3", "FY25 audited financial statements for Rice, Vanderbilt, Northwestern, Princeton, Harvard, Yale, Chicago, Duke, Emory, Penn, Stanford, Northeastern, BU, MIT (S&W, total operating expenses, depreciation, amortization, interest)", "Each university's finance / controller site", "e.g. bu.edu, web.mit.edu (vpf), finance.northeastern.edu, finance.harvard.edu, ...", "Blocked"),
    ("3", "IPEDS fall enrollment (EF2024) for the same 14 institutions", "nces.ed.gov/ipeds/datacenter", "nces.ed.gov", "Blocked"),
]
for i, row in enumerate(need, 5):
    for j, v in enumerate(row, 1):
        c = ws.cell(i, j, v); c.font = font(); c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(i, 5).font = font(bold=True, color=RED)

wb.save(OUT)
print("saved", OUT)
