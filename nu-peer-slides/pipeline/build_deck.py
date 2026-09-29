"""Assemble the 3-slide Northeastern-branded deck (interim build on the data in the original file)."""
import json, sys
import numpy as np
import pandas as pd
import openpyxl
from nu_helpers import *
from slide1 import build_slide1
from slide2 import build_slide2
from slide3 import build_slide3

OUT = sys.argv[1] if len(sys.argv) > 1 else "deck.pptx"
DRAFT = True

# ---------------------------------------------------------------- slide 1 data
df = pd.read_csv("inputs/slide1_scorecard_oct2023.csv")
df.loc[(df.name == "St. John's College") & (df.st == "NM"), "name"] = "St. John's College (Santa Fe)"
lm = json.load(open("inputs/s1_label_map_fixed.json"))
lm["Boston University"] = ["Boston University", 26170, 80582, 2]
ranked = df.sort_values("y", ascending=False).reset_index(drop=True)
foc = "Northeastern University"
foc_rank = int(ranked.index[ranked.name == foc][0]) + 1
fr = df[df.name == foc].iloc[0]
vu_rank = int(ranked.index[ranked.name == "Vanderbilt University"][0]) + 1
n_all = len(df)

s1_notes = [
    "(1) Median earnings of federally aided students who enroll each year and are employed but not enrolled, measured in the 10th year after enrollment.",
    "(2) Average net price for undergraduate Title IV-receiving students: full cost of attendance (tuition and fees, books and supplies, living expenses) minus federal, state and institutional aid.",
    "(3) Landmark College is located in Putney, Vermont and is dedicated to students with learning disabilities (ADHD, dyslexia and autism).",
    "Note: Of 6,543 IPEDS-reporting institutions, the analysis includes 1,557 (587 public, 970 private) with CONTROL = public or private not-for-profit; CCBASIC = 15–23; CCUGPROF = 5–15; NPT4 not null; and MD_EARN_WNE_P10 reported.",
]
s1_source = ("Source: College Scorecard (collegescorecard.ed.gov/data), data last updated October 10, 2023 "
             "(includes data from 1996 through 2022 for all undergraduate degree-granting institutions); Northeastern analysis.")
s1_speaker = f"""WHAT CHANGED FROM THE ORIGINAL
- Rebuilt on the Northeastern template as a native chart (data behind it is editable: right-click > Edit Data; institution names are in column C of the embedded sheet).
- Vanderbilt logo removed; Vanderbilt is now a navy label like the other private institutions. Northeastern is the red marker and label.
- Axis, title and footnote 1 now say "after entry": the Scorecard field MD_EARN_WNE_P10 is measured 10 years after enrollment, not 10 years after graduation (the original title/axis said graduation while its footnote said enrollment).
- Removed Vanderbilt-internal speaker notes and the 'VU Factbook' reference.

CHECKS RUN ON THE ORIGINAL DATA (Oct 2023 Scorecard release, 1,557 institutions)
- 587 public + 970 private = 1,557. Vanderbilt ranks {vu_rank}th by earnings (matches the original title; no ties). Northeastern ranks {foc_rank}th ({fr.y:,.0f} earnings; net price {fr.x:,.0f}).
- The plotted trendline is an ordinary least-squares fit on private institutions only (slope 1.035, intercept $29,753, r = 0.49, R-squared about 0.24): a weak fit, so 'better deal / worse deal' is a directional read, not a finding.

STATUS: prior-vintage data. Refresh with the latest College Scorecard institution file and re-run the same filters (see Note). University of the Sciences (merged into Saint Joseph's in 2022) should drop out on refresh.
"""

# ---------------------------------------------------------------- slide 2 data
raw = json.load(open("inputs/slide2_data_orig.json"))
s2 = {k: {int(y): v for y, v in d.items()} for k, d in raw.items()}
years = list(range(2015, 2025))            # 2014 removed per request
s2_speaker = """WHAT CHANGED FROM THE ORIGINAL
- 2014 row removed as requested. Rebuilt as a native chart; values are in the embedded sheet (x = percent agreeing, y = survey year).
- Colors: Faculty red, Trustees grey, Sr. Administrators black (original used yellow / white / blue letter badges).
- The source line now reads 2015-2024 and the 2013 trustee note was dropped because 2013 is no longer shown.
- Wording added to the scale note: respondents are business officers rating each group. These are perceptions of others, not self-reports by faculty or trustees, and business officers are themselves part of the 'Sr. Administrators' group.

CHECKS
- All 30 plotted values (2015-2024 x 3 groups) were read from the original slide's labels. The original hand-placed dots sat within 1.2 points of their labels on average (max 3.6), so the rebuild changes no reading.
- Gap between Faculty and Trustees: 46 points in 2015 (33% vs 79%), 47 points in 2024 (44% vs 91%). The gap is persistent rather than widening; Sr. Administrators vs Faculty narrowed from 55 to 51 points.

STATUS: 2025 (and a newer 2026 edition published July 2026) not yet added. The question wording has to be confirmed in those reports before plotting; see cover note.
"""

# ---------------------------------------------------------------- slide 3 data (FY23 from the original workbook)
wb = openpyxl.load_workbook("inputs/orig_slide3_workbook.xlsx", data_only=True)
ws = wb.active
rows3 = []
for r in range(94, 105):
    nm, sw, ox, st = ws.cell(r, 13).value, ws.cell(r, 14).value, ws.cell(r, 15).value, ws.cell(r, 16).value
    nm = {"U Chicago": "Chicago", "U Penn": "Penn"}.get(nm, nm)
    rows3.append(dict(name=nm, sw=sw, opex=ox, students=st))
HEALTH = ("Chicago", "Emory", "Penn", "Duke", "Stanford")
s3_notes = ["(1) Chicago, Emory, Penn, Duke and Stanford include health care systems.",
            "Note: Data from university FY23 reports. S&W = salaries and wages. Core operating expenses exclude depreciation, amortization and interest."]
s3_source = "Source: University FY23 financial reports; Northeastern analysis."
s3_speaker = """WHAT CHANGED FROM THE ORIGINAL
- University logos removed. In the original they were hand-placed pictures floating over a chart that plotted only the eight iso-lines, so nothing tied a logo to its data. Every institution is now a real data point (x = core operating expense per student, y = S&W share) with a label.
- Original labels dropped the K on several values (Emory '$269', iso-lines '$160'...'$320') and truncated Emory's students ('16,00'); fixed. Stanford's x-value ($915,817) sat past the original axis maximum ($900,000); axis now runs to $1,000K.
- Title reworded to a descriptive FY23 statement; the original 'more cost-efficient than peers' claim is not supported by S&W per student alone (health systems and research enterprises inflate it) and will be reframed once Northeastern, BU and MIT are added.

CHECKS
- For all 11 institutions: S&W per student = S&W / students, opex per student = core opex / students, and S&W share = S&W / core opex reproduce the workbook. S&W per student = x-value times y-value, so each point sits on its iso-line.
- Emory (16,000) and Northwestern (23,000) student counts are round numbers, so they are estimates; the other counts look like exact enrollment. Refresh should use one consistent IPEDS fall-enrollment definition for every school.

STATUS: FY23 data. FY25 update and the Northeastern, Boston University and MIT rows are pending. MIT's figures include Lincoln Laboratory; Stanford's include SLAC and the hospitals; both need a footnote when added.
"""


def ordinal(n):
    return f"{n}{'th' if 10 <= n % 100 <= 20 else {1:'st',2:'nd',3:'rd'}.get(n % 10,'th')}"

s1_sub = (f"Northeastern ranks {ordinal(foc_rank)} (of {n_all:,}) in median earnings 10 years after entry: "
          f"${fr.y/1000:,.0f}K at a ${fr.x/1000:,.0f}K average net price")

prs = open_base("base.pptx")
tag1 = "DRAFT · Oct 2023 Scorecard data; refresh pending" if DRAFT else None
tag2 = "DRAFT · 2025 survey year pending" if DRAFT else None
tag3 = "DRAFT · FY23 data; FY25 update and Northeastern, BU, MIT pending" if DRAFT else None

build_slide1(prs, df, lm, focus=foc, title="Wide variation in ‘return’ on college investment", subtitle=s1_sub,
             notes_paras=s1_notes, source=s1_source, draft_tag=tag1, notes_text=s1_speaker)
build_slide2(prs, s2, years, title="Is there a dangerous disconnect in perception at our institutions?",
             subtitle="From Inside Higher Ed’s annual survey of college and university business officers",
             notes_paras=["Note: Ratings are business officers’ perceptions of each group, not self-reports by faculty or trustees."],
             source="Source: Inside Higher Ed, Survey of College and University Business Officers, 2015–2024.",
             draft_tag=tag2, notes_text=s2_speaker, frame_bottom=6.5)
build_slide3(prs, rows3, "FY23", title="S&W per student ranges from $51K (Rice, Vanderbilt) to $405K (Stanford)",
             subtitle="Core operating expense per student and S&W share of expense; dashed lines mark constant S&W per student",
             notes_paras=s3_notes, source=s3_source, draft_tag=tag3, notes_text=s3_speaker, health_system=HEALTH, frame_bottom=6.38)

prs.core_properties.title = "Peer comparison: return on investment, perceptions and salaries per student"
import datetime
cp = prs.core_properties
cp.author = "Northeastern University"; cp.last_modified_by = "Northeastern University"
cp.created = cp.modified = datetime.datetime(2026, 9, 29, 12, 0, 0); cp.revision = 1; cp.comments = ""
prs.save(OUT)
print("saved", OUT, "focus rank", foc_rank, "vu rank", vu_rank)
