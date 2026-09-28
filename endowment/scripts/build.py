#!/usr/bin/env python3
"""Build the FY26 endowment returns deck and workbook from endowment/data/top50.json.

Outputs (endowment/output/):
  FY26_Endowment_Returns_Top50.pptx  5 slides: 1-yr, 5-yr, 10-yr return charts, then two
                                     static FY25 backup slides (assets/backup_render/*.jpg)
  FY26_Endowment_Data_Top50.xlsx     data sheet + methodology sheet

Usage: python3 endowment/scripts/build.py
Requires: python-pptx, openpyxl
"""
import json
import math
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data" / "top50.json"
OUT = ROOT / "output"
ASSETS = ROOT / "assets" / "backup_render"
DECK = OUT / "FY26_Endowment_Returns_Top50.pptx"
BOOK = OUT / "FY26_Endowment_Data_Top50.xlsx"

HOME = "Northeastern"
FONT = "Lato"
INK = RGBColor(0x1E, 0x20, 0x28)
GREY_TXT = RGBColor(0x6E, 0x6F, 0x73)
BAR = RGBColor(0xD6, 0xD7, 0xD9)
RED = RGBColor(0xC8, 0x10, 0x2E)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
GRID = RGBColor(0xE5, 0xE5, 0xEA)

METRICS = [
    ("fy26Return", "1-Year Return", "this figure"),
    ("fy26Return5yr", "5-Year Annualized Return", "this metric"),
    ("fy26Return10yr", "10-Year Annualized Return", "this metric"),
]
TITLE = "FY26 Endowment Returns — Top 50 National Universities"
RANKING = "2027 Best National Universities (released Sept 22, 2026)"


def load():
    d = json.loads(DATA.read_text())
    return d, d["institutions"]


def fmt_b(m):
    return "NA" if m is None else f"${m / 1000:,.1f}B"


# ---------------------------------------------------------------- deck
def text(slide, x, y, w, h, s, size, color=INK, bold=False, align=PP_ALIGN.LEFT, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = s
    f = r.font
    f.name, f.size, f.bold, f.italic = FONT, Pt(size), bold, italic
    f.color.rgb = color
    return tb


def swatch(slide, x, y, color):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.2), Inches(0.2))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.color.rgb = RGBColor(0, 0, 0)
    sh.line.width = Pt(0.5)
    sh.shadow.inherit = False


def footer(slide, s):
    text(slide, 0.4, 6.88, 10.6, 0.55, s, 8.5, GREY_TXT)
    text(slide, 11.1, 6.88, 1.83, 0.25, "Northeastern University", 9, INK, bold=True, align=PP_ALIGN.RIGHT)


def sizes(n):
    if n <= 10:
        return 14, 12, True
    if n <= 20:
        return 11, 10, True
    if n <= 32:
        return 9, 8, False
    return 7, 7, False


def label_every_bar(axis, vertical):
    """Stop PowerPoint from skipping category labels; rotate them when bars are narrow."""
    ax = axis._element
    for tag in ("c:tickLblSkip", "c:tickMarkSkip"):
        el = ax.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            ax.append(el)
        el.set("val", "1")
    # schema order: tickLblSkip and tickMarkSkip follow c:lblOffset and precede c:noMultiLvlLbl
    nm = ax.find(qn("c:noMultiLvlLbl"))
    if nm is not None:
        ax.remove(nm)
        ax.append(nm)
    if vertical:
        body = ax.find(qn("c:txPr")).find(qn("a:bodyPr"))
        body.set("rot", "-5400000")
        body.set("vert", "horz")


def chart_slide(prs, rows, key, label, noun, as_of, n_total):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, 0.4, 0.28, 12.53, 0.4, TITLE, 22, bold=True)
    text(s, 0.4, 0.68, 12.53, 0.3, label, 15, bold=True)
    rep = sorted([r for r in rows if r.get(key) is not None], key=lambda r: (-r[key], r["short"]))
    n = len(rep)
    if n == 0:
        text(s, 0.4, 1.1, 12.53, 1.0,
             f"As of {as_of}, no institution in the {RANKING} top 50 has publicly reported {noun} for FY26 yet. "
             "This slide will populate automatically as schools release FY26 investment results.",
             15, GREY_TXT)
        footer(s, "Source: institution financial reports, board/trustee disclosures, and financial press coverage.")
        return
    spacex = key == "fy26Return" and any(r.get("spacexPts") for r in rep)
    desc = f"As officially reported (updated {as_of}) — {n} of {n_total} institutions had released {noun}."
    if spacex:
        desc += (" Gold portion of a bar is the estimated share of the return attributable to legacy SpaceX "
                 "venture stakes crystallized by SpaceX's 2026 IPO.")
    text(s, 0.4, 1.02, 12.53, 0.5, desc, 13, GREY_TXT)

    lbl_pt, cat_pt, show_value = sizes(n)
    cd = CategoryChartData()
    stacked = n <= 20
    cd.categories = [f"{r['short']}\n(#{r['usNewsRank']})" if stacked else f"{r['short']} (#{r['usNewsRank']})"
                     for r in rep]
    totals = [r[key] for r in rep]
    base = [r[key] - ((r.get("spacexPts") or 0) if spacex else 0) for r in rep]
    cd.add_series("Total", totals)
    cd.add_series("Base return", base)
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.4), Inches(1.6),
                            Inches(12.53), Inches(5.0), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.font.name = FONT
    ch.font.size = Pt(cat_pt)
    ch.font.color.rgb = INK
    plot = ch.plots[0]
    plot.overlap = 100
    plot.gap_width = 45 if n <= 20 else 30
    tot_s, base_s = plot.series
    for i, r in enumerate(rep):
        home = r["short"] == HOME
        for ser, color in ((tot_s, GOLD if (spacex and r.get("spacexPts")) else BAR), (base_s, BAR)):
            pt = ser.points[i]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = RED if home else color
            pt.format.line.color.rgb = RGBColor(0x8A, 0x8B, 0x8E)
            pt.format.line.width = Pt(0.5)
        dl = tot_s.points[i].data_label
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        tf = dl.text_frame
        suffix = "%" if n <= 32 else ""
        tf.text = f"{r[key]:.1f}{suffix}"
        run = tf.paragraphs[0].runs[0]
        run.font.size, run.font.bold, run.font.name = Pt(lbl_pt), True, FONT
        run.font.color.rgb = RED if home else INK
        if show_value and key == "fy26Return":
            p2 = tf.add_paragraph()
            r2 = p2.add_run()
            r2.text = fmt_b(r.get("fy26M"))
            r2.font.size, r2.font.name = Pt(max(lbl_pt - 3, 7)), FONT
            r2.font.color.rgb = GREY_TXT
    va = ch.value_axis
    hi = max(totals)
    lo = min(0, min(totals))
    va.maximum_scale = (math.floor(hi / 5) + 1) * 5 + (5 if show_value else 0)
    va.minimum_scale = math.floor(lo / 5) * 5
    va.major_unit = 5
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = GRID
    va.format.line.fill.background()
    va.tick_labels.number_format = '0"%"'
    va.tick_labels.number_format_is_linked = False
    va.tick_labels.font.size = Pt(11)
    ca = ch.category_axis
    ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    ca.format.line.color.rgb = RGBColor(0, 0, 0)
    ca.tick_labels.font.size = Pt(cat_pt)
    label_every_bar(ca, vertical=not stacked)

    x = 0.4
    swatch(s, x, 6.66, BAR)
    text(s, x + 0.28, 6.64, 1.4, 0.25, "Base return", 11)
    x += 1.6
    if spacex:
        swatch(s, x, 6.66, GOLD)
        text(s, x + 0.28, 6.64, 2.6, 0.25, "Est. SpaceX-attributable gain", 11)
        x += 2.9
    swatch(s, x, 6.66, RED)
    text(s, x + 0.28, 6.64, 1.6, 0.25, "Northeastern", 11)

    left = n_total - n
    src = (f"Source: institution financial reports, board/trustee disclosures, and financial press coverage. "
           f"US News rank in parentheses ({RANKING}; ties included).")
    if show_value and key == "fy26Return":
        src += " Dollar figure under each return is the reported FY26 endowment value; NA = not yet published."
    if any(r["short"] == HOME and r.get("returnBasis") == "management estimate" for r in rep):
        src += " Northeastern's return is a management estimate pending audited figures."
    if spacex:
        src += " SpaceX-attributable shares reflect each institution's own disclosure; not an audited breakdown."
    src += f" Remaining {left} of {n_total} institutions had not yet released this figure as of {as_of}."
    footer(s, src)


def backup_slide(prs, img, title):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, 0.4, 0.12, 12.53, 0.4, title, 14, GREY_TXT, bold=True)
    h = 6.75
    w = h * 2666 / 1500
    s.shapes.add_picture(str(img), Inches((13.333 - w) / 2), Inches(0.55), Inches(w), Inches(h))


def build_deck(d, rows):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12188952), Emu(6858000)
    for key, label, noun in METRICS:
        chart_slide(prs, rows, key, label, noun, d["asOf"], len(rows))
    backup_slide(prs, ASSETS / "slide-1.jpg",
                 "BACKUP (for reference) — FY25 peer-cohort view, 1-Year Performance, n=123 (Charles Skorina & Co.)")
    backup_slide(prs, ASSETS / "slide-2.jpg",
                 "BACKUP (for reference) — FY25 peer-cohort view, 10-Year Annualized Performance, n=123 (Charles Skorina & Co.)")
    prs.save(DECK)


# ---------------------------------------------------------------- workbook
def build_book(d, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "FY26 Endowment Data"
    cols = [("US News Rank", 10), ("Institution", 36), ("Control", 9), ("FY End", 8),
            ("FY26 1-Yr Return %", 12), ("FY26 5-Yr Ann. %", 12), ("FY26 10-Yr Ann. %", 12),
            ("Est. SpaceX pts", 10), ("FY25 Endowment ($M)", 14), ("FY26 Endowment ($M)", 14),
            ("Last Checked", 12), ("Notes / Sources", 80)]
    hdr_font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="000000")
    rule = Side(style="thin", color="D9D9D9")
    for i, (h, w) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font, c.fill = hdr_font, hdr_fill
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center" if i != 2 and i != 12 else "left")
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[1].height = 32
    na = lambda v: "NA" if v is None else v
    for r_i, r in enumerate(sorted(rows, key=lambda r: (r["usNewsRank"], r["name"])), 2):
        vals = [r["usNewsRank"], r["name"], r["control"].title(), r["fyEnd"],
                na(r.get("fy26Return")), na(r.get("fy26Return5yr")), na(r.get("fy26Return10yr")),
                r.get("spacexPts") if r.get("spacexPts") is not None else "",
                na(r.get("fy25M")), na(r.get("fy26M")), r.get("lastChecked", ""), r["notes"]]
        home = r["short"] == HOME
        for c_i, v in enumerate(vals, 1):
            c = ws.cell(row=r_i, column=c_i, value=v)
            c.font = Font(name=FONT, size=10.5, bold=home, color="C8102E" if home else "000000")
            c.border = Border(bottom=rule)
            if c_i in (9, 10) and isinstance(v, (int, float)):
                c.number_format = "#,##0"
            elif c_i in (5, 6, 7, 8) and isinstance(v, (int, float)):
                c.number_format = "0.0"
            c.alignment = Alignment(horizontal="left" if c_i in (2, 12) else "right" if c_i >= 5 else "center",
                                    vertical="top", wrap_text=c_i == 12)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(cols)).column_letter}{len(rows) + 1}"

    m = wb.create_sheet("Methodology")
    m.column_dimensions["A"].width = 130
    rep = {k: [r for r in rows if r.get(k) is not None] for k, _, _ in METRICS}
    one = sorted(rep["fy26Return"], key=lambda r: -r["fy26Return"])
    listing = ", ".join(
        f"{r['name']} ({r['fy26Return']:.1f}%{', management estimate' if r.get('returnBasis') == 'management estimate' else ''})"
        for r in one)
    lines = [
        "Methodology & Caveats",
        f"Universe: {d['universe']}",
        "FY26 = fiscal year ending on or about June 30, 2026; several schools use a different fiscal year end "
        "(shown in the FY End column as MM-DD, e.g. Stanford 08-31, Caltech 09-30, UT Austin 08-31, Emory 08-31, Villanova 05-31).",
        f"As of {d['asOf']}: {len(one)} of {len(rows)} institutions had released an FY26 1-year return"
        + (f": {listing}." if one else ".")
        + f" 5-year annualized: {len(rep['fy26Return5yr'])} of {len(rows)}. 10-year annualized: {len(rep['fy26Return10yr'])} of {len(rows)}.",
        "NA = not yet published. Most institutions release audited endowment figures and returns between October and December; "
        "5- and 10-year annualized figures often follow later in investment-committee reports.",
        "Returns are taken directly from each institution's financial statements, board/trustee disclosures, investment office reports, "
        "or reputable press coverage quoting the institution. Returns are never calculated or derived from beginning/ending endowment values, "
        "and 1-, 5- and 10-year figures are recorded separately as reported.",
        "Dollar figures not verified against a primary source (audited statements, Form 990, official investment/foundation report) are flagged in Notes.",
        "Several public universities' endowments are held by affiliated foundations or investment companies (UC campus foundations, UVIMCO, UTIMCO, "
        "UNC Management Company, UFICO); the managing entity is noted where relevant.",
        "Est. SpaceX pts: percentage points of the 1-year return that a source explicitly attributes to legacy SpaceX stakes crystallized by SpaceX's 2026 IPO. "
        "Blank unless disclosed.",
        "Northeastern's FY26 return, where marked as a management estimate, is not independently verified or audited.",
    ]
    for i, s in enumerate(lines, 1):
        c = m.cell(row=i, column=1, value=s)
        c.font = Font(name=FONT, size=13 if i == 1 else 10.5, bold=i == 1)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(BOOK)
    return {k: len(v) for k, v in rep.items()}


def main():
    OUT.mkdir(exist_ok=True)
    d, rows = load()
    build_deck(d, rows)
    counts = build_book(d, rows)
    print(f"built {DECK.name} and {BOOK.name}: {len(rows)} institutions; reported "
          f"1-yr {counts['fy26Return']}, 5-yr {counts['fy26Return5yr']}, 10-yr {counts['fy26Return10yr']}")


if __name__ == "__main__":
    main()
